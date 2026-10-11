from pathlib import Path
import hashlib
import json
import re
import sys
import build_nh3_manuscript as M
import build_nh3_manuscript_v7 as M7
import nh3_docx as W
from build_nh3_manuscript import P, F, pr, pa, iv, K2, x3, cis, cin
ROOT, ST = (M7.ROOT, M7.ST)
OUT = ST / 'submission_ajpc' / 'manuscript'
TITLE = M7.TITLE
AUTHOR = 'Yunhao Jiang'
AFFIL = 'The University of Edinburgh, Old College, South Bridge, Edinburgh EH8 9YL, United Kingdom'
EMAIL = 'biojiang777@gmail.com'
ORCID = 'https://orcid.org/0009-0004-5954-7315'
KEYWORDS = ['PREVENT equations', 'Cardiovascular risk assessment', 'Measurement variability', 'Reproducibility', 'Statin eligibility', 'Hypertension guideline', 'Net benefit']

def _ai_declaration():
    p = ST / 'submission_ajpc' / 'declarations' / 'AI Use Declaration.docx'
    if not p.exists():
        return None
    from docx import Document
    paras = [q.text.strip() for q in Document(p).paragraphs if q.text.strip()]
    body = [t for t in paras if not t.lower().startswith('ai use declaration')]
    return ' '.join(body) if body else None
AI_DECLARATION = _ai_declaration()
V, e5, e3, e75, e30, e10, ETH, rec, lfs, ZTXT = (M7.V, M7.e5, M7.e3, M7.e75, M7.e30, M7.e10, M7.ETH, M7.rec, M7.lfs, M7.ZTXT)

def abstract():
    recs = sorted((float(rec(e)) for _, e, _, _ in ETH))
    lo4 = min((e['loss_fraction_single'] for e in (e3, e5, e75, e30)))
    hi4 = max((e['loss_fraction_single'] for e in (e3, e5, e75, e30)))
    return ['**Background:** PREVENT 10-year risk estimates enter the 2026 dyslipidemia and 2025 blood pressure guidelines, usually from one measurement set. We estimated how often the atherosclerotic cardiovascular disease (ASCVD) category and derived recommendations change on remeasurement, where, and at what expected cost.', f"**Methods:** In NHANES III, {pr['n']} adults aged 30–79 years without cardiovascular disease were reexamined under the same protocol a median of {iv['median']:.0f} days later, with all routine PREVENT inputs remeasured. The prespecified outcome was the proportion changing base PREVENT-ASCVD category. Post hoc analyses applied the 2025 and 2026 treatment rules and a decision-curve measurement-error model; NHANES 2009–2010 (n = {pa['n']:,}) isolated the urine specimen effect.", f"**Results:** {pr['events']} participants ({cis(pr)}) changed category; the chance was {100 * min(M.pi0):.0f}–{100 * max(M.pi0):.0f}% at a threshold and ≥25% only within about {ZTXT} risk. Blood pressure and lipids contributed {F(K2['Psi']['BP']['pp'])} and {F(K2['Psi']['LIPID']['pp'])} percentage points. Single-examination recommendations changed for {P(x3['theta_rx']['estimate'])}% of {x3['n_untreated_eligible']} untreated participants (2025 blood pressure rule) and {P(V['primary_recommended_change']['estimate'])}% of {V['n']} statin-naive participants (2026 Class 1 statin rule). Single-visit decisions forfeited {P(lo4)}–{P(hi4)}% of the achievable net benefit at the 3%, 5%, 7.5% and 30-year 10% thresholds and {lfs(e10)} at 10%; remeasurement within the zones recovered {recs[0]:.0f}–{recs[-1]:.0f}% of the loss. The specimen change altered {P(pa['estimate'])}% of urine albumin-to-creatinine ratio model categories.", '**Conclusions:** One in 10 adults changed PREVENT category on remeasurement, within narrow zones around the thresholds, and the changes reached current treatment rules; the expected net benefit at stake is small, the cost is inconsistent advice, which targeted remeasurement can limit.']

def methods_ajpc():
    out = []
    for b in M7.methods():
        if b[0] == 'p' and 'this analysis used deidentified public data.' in b[1]:
            out.append(('p', b[1].replace('this analysis used deidentified public data.', 'this secondary analysis of publicly available, deidentified data did not require institutional review board review. Reporting follows the STROBE recommendations for observational studies.')))
        else:
            out.append(b)
    return out
JREF = re.compile('^(?P<authors>[^.]+?)\\. (?P<title>.+?)\\. (?P<journal>[A-Z][^.;]*?)\\.? (?P<year>(?:19|20)\\d{2});(?P<vol>\\d+)(?:\\((?P<issue>[^)]*)\\))?:(?P<pages>[^.\\s]+)\\.?(?: (?P<doi>https?://\\S+))?$')

def els_author(a):
    a = a.strip()
    if a in ('et al', 'et al.'):
        return 'et al.'
    m = re.match('^(.+?)\\s([A-Z]{1,3})$', a)
    return '.'.join(m.group(2)) + '. ' + m.group(1) if m else a
JREF_NOPAGES = re.compile('^(?P<authors>[^.]+?)\\. (?P<title>.+?)\\. (?P<journal>[A-Z][^.;]*?)\\.? (?P<year>(?:19|20)\\d{2});(?P<vol>\\d+)(?:\\((?P<issue>[^)]*)\\))?\\.?(?: (?P<doi>https?://\\S+))?$')
WREF = re.compile('^(?P<auth>[^.]+)\\. (?P<title>.+?)\\.? (?P<url>https?://\\S+?)\\. Accessed (?P<mon>[A-Z][a-z]+) (?P<day>\\d+), (?P<year>\\d{4})\\.$')
DREF = re.compile('^(?P<auth>[^.]+)\\. (?P<title>.+?), version (?P<ver>[\\d.]+)\\. Zenodo; (?P<year>\\d{4})\\. (?P<doi>https?://\\S+)$')

def elsevier_ref(s):
    s = s.strip()
    m = JREF.match(s)
    if m:
        g = m.groupdict()
        authors = ', '.join((els_author(a) for a in g['authors'].split(', ')))
        return f"{authors}, {g['title']}, {g['journal']} {g['vol']} ({g['year']}) {g['pages'].replace('-', '–')}." + (f" {g['doi']}" if g['doi'] else '')
    m = JREF_NOPAGES.match(s)
    if m:
        g = m.groupdict()
        authors = ', '.join((els_author(a) for a in g['authors'].split(', ')))
        return f"{authors}, {g['title']}, {g['journal']} {g['vol']} ({g['year']})." + (f" {g['doi']}" if g['doi'] else '')
    m = WREF.match(s)
    if m:
        g = m.groupdict()
        return f"{g['auth']}, {g['title']}. {g['url']}, {g['year']} (accessed {int(g['day'])} {g['mon']} {g['year']})."
    m = DREF.match(s)
    if m:
        g = m.groupdict()
        return f"[dataset] {els_author(g['auth'])}, {g['title']} [dataset], Zenodo, v{g['ver']}, {g['year']}. {g['doi']}"
    return s

def back_matter():
    return [('h1', 'CRediT authorship contribution statement'), ('p', f'{AUTHOR}: Conceptualization, Methodology, Software, Formal analysis, Investigation, Data curation, Validation, Visualization, Writing – original draft, Writing – review & editing, Project administration.'), ('h1', 'Declaration of competing interest'), ('p', 'The author declares no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper.'), ('h1', 'Funding'), ('p', 'This research did not receive any specific grant from funding agencies in the public, commercial, or not-for-profit sectors.'), ('h1', 'Data availability'), ('p', 'NHANES data and documentation are publicly available from the National Center for Health Statistics.[[nhanes3_files,nh_albcr_f,nh_l16_2_b]] Processed data and analysis code for Study 2 are archived on Zenodo (version 2.0.0).[[paper1]] Analysis code, protocols and derived data for Study 1 and the post hoc analyses are released as version 3.0.0 of the same repository (https://github.com/Re1nhard-1/uacr-prevent; Zenodo DOI: [to be inserted when the version is released]).'), ('h1', 'Supplementary material'), ('p', 'Supplementary material associated with this article (Supplemental Methods; Tables S1–S41; Figures S1–S9; Supplemental References; the frozen Study 1 protocol and the dated post hoc protocols with addenda) is provided as separate files.')] + ([('h1', 'Declaration of generative AI and AI-assisted technologies in the writing process'), ('p', AI_DECLARATION)] if AI_DECLARATION else [])

def legends():
    L = M7.legends()
    ci = 'Central Illustration. Remeasuring every PREVENT input about 2 weeks later changes the 10-year ASCVD risk category of about 1 in 10 adults, almost only within narrow zones around the 3%, 5% and 10% thresholds (left; NHANES III, n = 759; shading, zones where the modelled probability of a different category is ≥25%), while the expected net benefit forfeited by single-visit decisions is small at most guideline thresholds (right; measurement-error model; percent of the net benefit achievable by deciding on the short-term mean level; 95% percentile intervals from 1,000 bootstrap replicates). Full detail in Figures 1 and 4.'
    return [ci] + L

def build(check=False):
    abs_ = abstract()
    tp = [('title', TITLE), ('p', f'{AUTHOR}ᵃ'), ('p', f'ᵃ {AFFIL}'), ('p', f'Corresponding author: {AUTHOR}, {AFFIL}. E-mail: {EMAIL}. ORCID: {ORCID}'), ('p', 'Keywords: ' + '; '.join(KEYWORDS))]
    main = [('h1', 'Abstract')] + [('p', s) for s in abs_] + [('h1', '1. Introduction')] + [('p', s) for s in M7.intro()] + [('h1', '2. Methods')] + methods_ajpc() + [('h1', '3. Results')] + M7.results() + [('h1', '4. Discussion')] + M7.discussion() + back_matter()
    t1a, t1b = M.table1()
    t1a['title'] = 'Table 1. Study 1: Characteristics of Participants and Repeatability of Inputs'
    t1b['title'] = 'Table 1 (continued). B. Inputs at Each Examination'
    tables = [('break',), ('table', t1a), ('break',), ('table', t1b), ('break',), ('table', M7.table2())]
    figs = [('break',), ('h1', 'Figure captions')] + [('legend', s) for s in legends()]
    all_blocks = tp + main + tables + figs
    order = W.number_citations(all_blocks)
    text_blocks = [('p', s) for s in M7.intro()] + methods_ajpc() + M7.results() + M7.discussion()
    stats = {'abstract_words': M7.para_words(abs_), 'main_text_words': W.words(text_blocks), 'references': len(order), 'tables': 2, 'figures': 4, 'central_illustration': 1, 'keywords': len(KEYWORDS), 'limits': {'venue': 'American Journal of Preventive Cardiology (original research)', 'abstract_max': 250, 'text_usual_max': 4000, 'keywords_max': 7}}
    print(json.dumps(stats))
    OUT.mkdir(parents=True, exist_ok=True)
    refs_raw = [M7.ref_text(k) for k in sorted(order, key=order.get)]
    refs = [elsevier_ref(s) for s in refs_raw]
    if check:
        txt = [W.render_cites(x[1], order) for x in tp + main if x[0] in ('p', 'h1', 'h2', 'title')]
        (OUT / 'draft_text_check.txt').write_text('\n\n'.join(txt) + '\n\nREFERENCES\n' + '\n'.join((f'[{i}] {r}' for i, r in enumerate(refs, 1))), encoding='utf-8')
        return stats
    doc = W.new_doc(TITLE)
    W.render(doc, [('title', TITLE)], order)
    for s in [f'{AUTHOR}ᵃ', f'ᵃ {AFFIL}', f'Corresponding author: {AUTHOR}. E-mail: {EMAIL}. ORCID: {ORCID}', f"Word count: abstract {stats['abstract_words']}; main text {stats['main_text_words']} (Introduction to Discussion, excluding references, tables and captions). Tables 2; figures 4 plus a Central Illustration; references {stats['references']}.", 'Keywords: ' + '; '.join(KEYWORDS)]:
        W.para(doc, s)
    W.page_break(doc)
    W.render(doc, main, order)
    W.para(doc, 'References', style='Heading 1')
    W.render(doc, [('refs', refs)], order)
    W.render(doc, tables + figs, order)
    W.add_line_numbers(doc)
    out = OUT / 'nh3_manuscript_ajpc.docx'
    if out.exists():
        out = OUT / 'nh3_manuscript_ajpc_REBUILD_do_not_submit.docx'
        print('AUTHOR-EDITED MANUSCRIPT PRESERVED; rebuilt copy written to', out.name)
    doc.save(out)
    rec_ = {'venue': 'American Journal of Preventive Cardiology', 'built_from_revision': '7 draft 2 (build_nh3_manuscript_v7.py)', 'stats': stats, 'reference_order': sorted(order, key=order.get), 'references_elsevier_style': refs, 'placeholders': ['affiliation postal address', 'ORCID', 'Zenodo DOI for version 3.0.0', 'generative-AI declaration text'], 'source_hashes': {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in [ROOT / 'scripts/build_nh3_manuscript_v7.py', ST / 'EVIDENCE_LEDGER.md']}}
    (OUT / 'build_record_ajpc.json').write_text(json.dumps(rec_, indent=1) + '\n', encoding='utf-8')
    print('saved', out)
if __name__ == '__main__':
    build(check='--check' in sys.argv)

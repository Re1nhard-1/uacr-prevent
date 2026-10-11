import re
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, Mm
CITE = re.compile('\\[\\[([A-Za-z0-9_.,\\- ]+)\\]\\]')

def number_citations(blocks):
    order = {}

    def scan(s):
        for m in CITE.finditer(s or ''):
            for k in [x.strip() for x in m.group(1).split(',')]:
                if k not in order:
                    order[k] = len(order) + 1
    for b in blocks:
        kind = b[0]
        if kind in ('h1', 'h2', 'p', 'pc', 'small', 'title', 'legend'):
            scan(b[1])
        elif kind == 'table':
            spec = b[1]
            scan(spec.get('title', ''))
            scan(spec.get('note', ''))
            for r in spec['rows']:
                for c in r:
                    scan(c if isinstance(c, str) else '')
        elif kind == 'figure':
            scan(b[2])
    return order

def fmt_nums(ns):
    ns = sorted(set(ns))
    out = []
    i = 0
    while i < len(ns):
        j = i
        while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
            j += 1
        out.append(f'{ns[i]}–{ns[j]}' if j - i >= 2 else f'{ns[i]},{ns[j]}' if j == i + 1 else f'{ns[i]}')
        i = j + 1
    return '[' + ','.join(out) + ']'

def render_cites(s, order, prefix=''):

    def rep(m):
        ks = [x.strip() for x in m.group(1).split(',')]
        nums = [order[k] for k in ks]
        t = fmt_nums(nums)
        return t if not prefix else t.replace('[', '[' + prefix).replace(',', ',' + prefix).replace('–', '–' + prefix)
    return CITE.sub(rep, s)

def new_doc(title, author='Yunhao Jiang'):
    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.subject = ''
    doc.core_properties.author = author
    doc.core_properties.comments = ''
    doc.core_properties.keywords = ''
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption']:
        st = doc.styles[name]
        st.font.name = 'Times New Roman'
        st.font.size = Pt(12)
        rf = st.element.get_or_add_rPr().get_or_add_rFonts()
        for k in ['ascii', 'hAnsi', 'cs', 'eastAsia']:
            rf.set(qn('w:' + k), 'Times New Roman')
        for k in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            rf.attrib.pop(qn('w:' + k), None)
        st.font.bold = name in ['Title', 'Heading 1', 'Heading 2']
        st.font.italic = name == 'Heading 2'
        st.font.color.rgb = None
        st.paragraph_format.space_before = Pt(12 if name.startswith('Heading') else 0)
        st.paragraph_format.space_after = Pt(6 if name.startswith('Heading') else 0)
    doc.styles['Title'].font.size = Pt(14)
    doc.styles['Normal'].paragraph_format.line_spacing = 2.0
    for sec in doc.sections:
        sec.page_width, sec.page_height = (Mm(210), Mm(297))
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    add_page_numbers(doc)
    return doc

def add_page_numbers(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = 1
    r = p.add_run()
    for tag, text in [('begin', None), (None, 'PAGE'), ('end', None)]:
        if tag:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), tag)
            r._r.append(e)
        else:
            e = OxmlElement('w:instrText')
            e.set(qn('xml:space'), 'preserve')
            e.text = text
            r._r.append(e)

def add_line_numbers(doc):
    for sec in doc.sections:
        e = OxmlElement('w:lnNumType')
        e.set(qn('w:countBy'), '1')
        e.set(qn('w:restart'), 'continuous')
        e.set(qn('w:distance'), '360')
        sec._sectPr.append(e)
SUB = re.compile('(?<![A-Za-z0-9])([θΔπΨΣσfDy])_(\\{[^}]+\\}|[A-Za-z0-9∈]+)')

def add_runs(p, s, size=None, bold=None, italic=None):
    for part in re.split('(\\*\\*[^*]+\\*\\*)', s):
        if not part:
            continue
        b = part.startswith('**') and part.endswith('**')
        txt = part[2:-2] if b else part
        pos = 0
        for m in list(SUB.finditer(txt)) + [None]:
            end = m.start(2) - 1 if m else len(txt)
            segs = [(txt[pos:end], False)]
            if m:
                segs.append((m.group(2).strip('{}'), True))
                pos = m.end()
            for t, sub in segs:
                if not t:
                    continue
                r = p.add_run(t)
                if size:
                    r.font.size = Pt(size)
                if bold is not None or b:
                    r.font.bold = True if b else bold
                if italic is not None:
                    r.font.italic = italic
                if sub:
                    r.font.subscript = True

def para(doc, s, style='Normal', size=None, bold=None, italic=None, single=False, align=None, keep=False):
    p = doc.add_paragraph(style=style)
    add_runs(p, s, size, bold, italic)
    if single:
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(4)
    if align is not None:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_with_next = True
    return p

def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def shade(cell, fill='FFFFFF', top=True, bottom=True):
    tc = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for edge in ['top', 'left', 'bottom', 'right']:
        e = OxmlElement('w:' + edge)
        if edge in ('left', 'right'):
            e.set(qn('w:val'), 'nil')
        else:
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), '4')
            e.set(qn('w:color'), 'BFBFBF')
        b.append(e)
    tc.append(b)
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), fill)
    tc.append(s)

def fill(cell, text, bold=False, size=8.5, align=None, italic=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    if align is not None:
        p.alignment = align
    add_runs(p, text, size, bold, italic)

def table(doc, spec, order, prefix=''):
    para(doc, render_cites(spec['title'], order, prefix), bold=True, single=True, keep=True)
    ncol = len(spec['header'])
    size = spec.get('size', 8.5)
    t = doc.add_table(rows=0, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    def add(cells, header=False, section=False):
        r = t.add_row()
        cs = r.cells
        if section:
            m = cs[0].merge(cs[-1])
            fill(m, render_cites(cells[0], order, prefix), bold=True, size=size)
            shade(m, 'EEF2F5')
            return
        for j, (c, v) in enumerate(zip(cs, cells)):
            fill(c, render_cites(str(v), order, prefix), bold=header, size=size, align=None if j == 0 else 1)
            shade(c, 'E7EEF3' if header else 'FFFFFF')
    add(spec['header'], header=True)
    for row in spec['rows']:
        add(row, section=len(row) == 1)
    w = spec.get('widths_in')
    if w:
        for row in t.rows:
            for c, wi in zip(row.cells, w):
                c.width = Inches(wi)
    tr = t.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    h = OxmlElement('w:tblHeader')
    h.set(qn('w:val'), 'true')
    trPr.append(h)
    if spec.get('note'):
        para(doc, render_cites(spec['note'], order, prefix), size=9, single=True)

def render(doc, blocks, order, figure_width_in=6.3, prefix=''):
    rc = lambda t: render_cites(t, order, prefix)
    for b in blocks:
        k = b[0]
        if k == 'title':
            para(doc, rc(b[1]), style='Title')
        elif k == 'h1':
            para(doc, b[1], style='Heading 1', keep=True)
        elif k == 'h2':
            para(doc, b[1], style='Heading 2', keep=True)
        elif k == 'p':
            para(doc, rc(b[1]))
        elif k == 'pc':
            para(doc, rc(b[1]), align=1)
        elif k == 'small':
            para(doc, rc(b[1]), size=10, single=True)
        elif k == 'legend':
            para(doc, rc(b[1]), single=True, size=11)
        elif k == 'table':
            table(doc, b[1], order, prefix)
        elif k == 'figure':
            doc.add_picture(str(b[1]), width=Inches(b[3] if len(b) > 3 else figure_width_in))
            doc.paragraphs[-1].alignment = 1
            para(doc, rc(b[2]), single=True, size=10)
        elif k == 'break':
            page_break(doc)
        elif k == 'refs':
            for i, s in enumerate(b[1], 1):
                para(doc, f"{(b[2] if len(b) > 2 else '')}{i}. {s}", single=True, size=10.5)
        else:
            raise ValueError(k)

def words(blocks):
    n = 0
    for b in blocks:
        if b[0] == 'p':
            n += len(re.findall('\\S+', CITE.sub('', b[1]).replace('**', '')))
    return n
